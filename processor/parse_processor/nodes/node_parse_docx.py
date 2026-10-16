'''processor/parse_processor/nodes/node_parse_docx.py'''

'''
步骤分解
1. 接收状态: 获取 `docx_path`（回退 stored_file_path / import_file_path）。
2. 打开文档: 用 python-docx 打开文件，失败统一转成 DocxParseError。
3. 按文档顺序遍历: 段落与表格必须按原文出现的顺序输出，不能分成两个列表各自遍历
              （否则条款顺序错乱，后面切分和定位全废）。
4. 转换为统一中间结构: 每条产出 { seq, page_no, paragraph_index, text, bbox, kind }。
5. 正文为空判定: 全文去空白后为空 → 报错（对应 SPEC 11.1「正文为空」）。
6. 写回状态 + 统计日志。

需求分析
输入:
- `docx_path`：DOCX 文件路径（node_detect_type 权威判定后写入）

输出:
- `fragments`：DocFragment[] 统一中间结构（段落 + 表格）

关于页码（重要）:
    python-docx **拿不到真实页码**——页码是排版结果，不是文档结构。
    所以 DOCX 来源的 page_no 统一记 0，表示「不适用」。
    前端工作台预览 DOCX 走 Mammoth.js 转 HTML，不依赖页码，只依赖 seq / paragraph_index，
    因此这个取舍不影响高亮联动；但 PDF 来源的 page_no 是真实的（node_parse_pdf 会逐页输出）。
'''

import json
import logging
from pathlib import Path
from typing import Iterator, List, Dict, Any

from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

from processor.parse_processor.base import BaseNode, setup_logging
from processor.parse_processor.exceptions import (
    DocxParseError,
    FileProcessingError,
    ParseProcessError,
    StateFieldError,
)
from processor.parse_processor.state import ParseGraphState, get_default_state


def iter_block_items(parent) -> Iterator[Any]:
    """
    按【文档顺序】迭代 DOCX 里的段落与表格。

    为什么需要这个函数:
        python-docx 的 document.paragraphs 与 document.tables 是两个彼此独立的列表，
        各自遍历会丢失两者的相对顺序（合同里表格往往夹在条款段落中间）。
        这里直接遍历 body 的 XML 子元素，才能拿到真实的文档顺序。
    """
    if isinstance(parent, DocxDocument):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise DocxParseError(message=f'不支持的 DOCX 父容器类型：{type(parent).__name__}')

    for child in parent_elm.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)


class NodeParseDocx(BaseNode):
    '''DOCX 解析节点：按文档顺序抽取段落与表格，输出统一中间结构 DocFragment[]'''
    name = 'node_parse_docx'

    # 块类型常量
    KIND_PARAGRAPH = 'paragraph'
    KIND_TABLE = 'table'

    def process(self, state: ParseGraphState) -> ParseGraphState:
        """
        节点：DOCX 解析（node_parse_docx）
        整体流程：取路径校验 → 打开文档 → 按顺序遍历块 → 转中间结构 → 正文为空判定 → 写回状态
        核心目的：把 DOCX 的「排版结构」拆成带序号的 DocFragment[]，
                 为后续条款切分与原文定位提供稳定的坐标基准
        必要参数：docx_path
        更新参数：fragments

        :param state: 工作流状态对象
        :return: 更新后的状态对象
        """
        # =======步骤一:获取并校验文件路径=======
        file_path = self._step_1_get_inputs(state)

        # =======步骤二:按文档顺序遍历段落与表格=======
        # 作用：核心抽取逻辑，段落与表格必须按原文顺序输出
        fragments = self._step_2_extract_fragments(file_path)

        # =======步骤三:正文为空判定=======
        self._step_3_check_content(fragments, file_path)

        # =======步骤四:写回状态 + 输出统计日志=======
        state['fragments'] = fragments
        self._step_4_print_stats(file_path, fragments)

        return state

    def _step_1_get_inputs(self, state: ParseGraphState) -> Path:
        """
        【步骤1】获取并校验文件路径
        功能：优先取 docx_path，回退 stored_file_path / import_file_path；校验存在性与是否为文件
        :param state: 项目状态字典（ParseGraphState）
        :return: 校验通过后的文件 Path 对象
        """
        file_path_str = (
            state.get('docx_path')
            or state.get('stored_file_path')
            or state.get('import_file_path')
        )
        if not file_path_str:
            raise StateFieldError(
                field_name='docx_path',
                message='DOCX 文件路径不能为空',
                expected_type=str
            )

        file_path = Path(file_path_str)
        if not file_path.exists():
            raise FileProcessingError(message=f'文件 {file_path.name} 不存在')
        if not file_path.is_file():
            raise FileProcessingError(message=f'路径 {file_path} 不是一个文件')

        return file_path

    def _step_2_extract_fragments(self, file_path: Path) -> List[Dict[str, Any]]:
        """
        【步骤2】按文档顺序抽取段落与表格
        功能：遍历文档块，段落取 text，表格按行拼成 " | " 分隔的文本；
             跳过纯空白块，但块序号照常递增，保证定位坐标稳定
        :param file_path: DOCX 文件 Path 对象
        :return: DocFragment[] 中间结构列表
        """
        try:
            import docx

            document = docx.Document(str(file_path))
        except PackageNotFoundError as e:
            raise DocxParseError(
                message=f'文件 {file_path.name} 不是有效的 DOCX 包（或已损坏）',
                cause=e
            )
        except Exception as e:
            raise DocxParseError(
                message=f'打开 DOCX 文件 {file_path.name} 失败：{e}',
                cause=e
            )

        fragments: List[Dict[str, Any]] = []
        seq = 0  # 有效块计数（写入 seq）
        block_index = 0  # 文档块序号（含被跳过的空白块，保证坐标稳定）

        for block in iter_block_items(document):
            if isinstance(block, Paragraph):
                text = block.text
                kind = self.KIND_PARAGRAPH
            elif isinstance(block, Table):
                text = self._table_to_text(block)
                kind = self.KIND_TABLE
            else:
                block_index += 1
                continue

            current_index = block_index
            block_index += 1

            # 跳过纯空白块（噪音），但序号位置已经占掉，不会导致后续块错位
            if not text.strip():
                continue

            fragments.append({
                'seq': seq,
                'page_no': 0,  # DOCX 无真实页码，记 0 表示不适用
                'paragraph_index': current_index,
                'text': text,
                'bbox': None,  # DOCX 无坐标
                'kind': kind,
            })
            seq += 1

        return fragments

    def _table_to_text(self, table: Table) -> str:
        """
        【辅助函数】把表格转成纯文本
        功能：逐行用 " | " 连接单元格文本，保留表格的行列可读性
        :param table: python-docx 的 Table 对象
        :return: 表格文本
        """
        rows_text = []
        for row in table.rows:
            cells_text = [cell.text.strip() for cell in row.cells]
            rows_text.append(' | '.join(cells_text))
        return '\n'.join(rows_text)

    def _step_3_check_content(self, fragments: List[Dict[str, Any]], file_path: Path) -> None:
        """
        【步骤3】正文为空判定
        功能：全文去空白后为空 → 解析不出任何内容，直接报错（对应 SPEC 11.1「正文为空」）
        :param fragments: DocFragment[] 列表
        :param file_path: 文件 Path 对象
        """
        full_text = ''.join(fragment['text'] for fragment in fragments)
        if not full_text.strip():
            raise DocxParseError(
                message=f'文件 {file_path.name} 正文为空，未解析出任何文本内容'
            )

    def _step_4_print_stats(self, file_path: Path, fragments: List[Dict[str, Any]]) -> None:
        """
        【步骤4】输出统计日志
        :param file_path: 文件 Path 对象
        :param fragments: DocFragment[] 列表
        """
        paragraph_count = sum(1 for f in fragments if f['kind'] == self.KIND_PARAGRAPH)
        table_count = sum(1 for f in fragments if f['kind'] == self.KIND_TABLE)
        total_chars = sum(len(f['text']) for f in fragments)

        self.log_step('文件', file_path.name)
        self.log_step('段落数', str(paragraph_count))
        self.log_step('表格数', str(table_count))
        self.log_step('总字符数', str(total_chars))
        if fragments:
            preview = fragments[0]['text'][:30].replace('\n', ' ')
            self.log_step('首块预览', preview)


if __name__ == '__main__':
    # 激活日志的全局配置
    setup_logging()

    import tempfile
    import docx

    # 临时根目录：由系统生成，测试结束自动随进程清理
    root_dir = Path(tempfile.mkdtemp(prefix='node_parse_docx_test_'))

    def _make_docx(name: str, paragraphs: List[str], table: List[List[str]] = None) -> Path:
        """用 python-docx 造一个真实的 DOCX（段落 + 可选表格）"""
        file_path = root_dir / name
        file_path.parent.mkdir(parents=True, exist_ok=True)

        document = docx.Document()
        for para_index, para_text in enumerate(paragraphs):
            document.add_paragraph(para_text)
            # 把表格插在第二段之后，用来验证段落/表格的文档顺序
            if table and para_index == 1:
                doc_table = document.add_table(rows=0, cols=len(table[0]))
                for row_data in table:
                    row_cells = doc_table.add_row().cells
                    for cell_index, cell_text in enumerate(row_data):
                        row_cells[cell_index].text = cell_text
        document.save(str(file_path))
        return file_path

    node_parse_docx = NodeParseDocx()
    passed = 0

    # ---------- 一、正常场景 ----------
    logging.getLogger().info('=' * 24 + ' 一、正常场景 ' + '=' * 24)

    sample = _make_docx(
        'sample/采购合同.docx',
        paragraphs=[
            '软件采购合同',
            '第一条 标的物',
            '第二条 付款方式',
            '第三条 违约责任',
        ],
        table=[
            ['条款', '金额'],
            ['首付款', '300000'],
        ],
    )

    init_state = get_default_state()
    init_state.update({
        'task_id': 'task-docx-ok',
        'docx_path': str(sample),
        'stored_file_path': str(sample),
        'is_docx_read_enabled': True,
    })
    result = node_parse_docx(init_state)
    fragments = result['fragments']

    assert len(fragments) == 5, f'应有 4 段落 + 1 表格 = 5 块，实际 {len(fragments)}'

    # 断言 1：seq 从 0 连续递增
    assert [f['seq'] for f in fragments] == list(range(len(fragments))), 'seq 应从 0 连续递增'

    # 断言 2：内容抽取正确（含表格 " | " 拼接）
    texts = [f['text'] for f in fragments]
    assert texts[0] == '软件采购合同', '第一块应为合同标题'
    assert texts[1] == '第一条 标的物', '第二块应为第一条'
    assert any('条款 | 金额' in t for t in texts), '表格应按 " | " 拼成文本'

    # 断言 3：段落与表格按文档顺序交织（表格插在第 2 段之后）
    table_seq = next(f['seq'] for f in fragments if f['kind'] == 'table')
    second_para_seq = next(f['seq'] for f in fragments if f['text'] == '第二条 付款方式')
    assert table_seq < second_para_seq, '表格应出现在第三段之前（验证文档顺序未被打乱）'

    # 断言 4：paragraph_index 与段落位置一致
    assert fragments[0]['paragraph_index'] == 0
    assert fragments[1]['paragraph_index'] == 1
    assert fragments[2]['paragraph_index'] == 2, '表格占掉 index=2，第三段应为 index=3 之前的位置'

    # 断言 5：DOCX 无页码/坐标，明确记 0 / None
    assert all(f['page_no'] == 0 for f in fragments), 'DOCX 来源 page_no 应统一为 0'
    assert all(f['bbox'] is None for f in fragments), 'DOCX 来源 bbox 应为 None'

    # 断言 6：kind 只含两种合法值
    assert {f['kind'] for f in fragments} <= {'paragraph', 'table'}, 'kind 取值非法'

    passed += 1
    logging.getLogger().info(f'  ✅ 段落+表格共 {len(fragments)} 块，顺序与坐标正确')

    # 纯段落文档
    plain = _make_docx('sample/纯段落.docx', paragraphs=['甲方：某某公司', '乙方：某某科技'])
    init_state = get_default_state()
    init_state['task_id'] = 'task-docx-plain'
    init_state['docx_path'] = str(plain)
    result = node_parse_docx(init_state)
    assert len(result['fragments']) == 2, '纯段落文档应解析出 2 块'
    assert all(f['kind'] == 'paragraph' for f in result['fragments'])
    passed += 1
    logging.getLogger().info('  ✅ 纯段落文档解析正确')

    # ---------- 二、异常场景 ----------
    logging.getLogger().info('=' * 24 + ' 二、异常场景 ' + '=' * 24)

    blank_docx = _make_docx('err/空白合同.docx', paragraphs=['', '   ', ''])
    fake_docx = root_dir / 'err' / '假文件.docx'
    fake_docx.write_bytes(b'this is definitely not a docx package')

    err_cases = [
        ('正文为空', {'docx_path': str(blank_docx)}, DocxParseError),
        ('非DOCX包', {'docx_path': str(fake_docx)}, DocxParseError),
        ('文件不存在', {'docx_path': str(root_dir / '不存在.docx')}, FileProcessingError),
        ('路径是目录', {'docx_path': str(root_dir)}, FileProcessingError),
        ('路径全为空', {}, StateFieldError),
    ]

    for case_name, overrides, expected_exc in err_cases:
        init_state = get_default_state()
        init_state['task_id'] = 'task-docx-error'
        init_state.update(overrides)
        try:
            node_parse_docx(init_state)
        except Exception as e:
            assert isinstance(e, ParseProcessError), f'{case_name} 应被基类包装成 ParseProcessError'
            assert isinstance(e.cause, expected_exc), \
                f'{case_name} 应抛出 {expected_exc.__name__}，实际为 {type(e.cause).__name__}'
            passed += 1
            logging.getLogger().info(f'  ✅ {case_name:<10} → 正确抛出 {expected_exc.__name__}')
        else:
            raise AssertionError(f'{case_name} 未抛出异常，用例失败')

    # ---------- 三、打印最终状态（便于人工核对）----------
    logging.getLogger().info('=' * 24 + ' 三、最终状态 ' + '=' * 24)
    logging.getLogger().info(json.dumps(result['fragments'], ensure_ascii=False, indent=4))

    logging.getLogger().info(f'node_parse_docx 自测全部通过 ✅ 共 {passed} 个用例')
