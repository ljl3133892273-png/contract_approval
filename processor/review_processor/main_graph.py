'''processor/review_processor/main_graph.py'''

'''
目标:
- 构建完整的审查工作流图
- 定义节点执行顺序
- 提供便捷的测试入口

需求分析
流程结构（纯顺序，无分支）:
1. `node_load_clause`   按 task_id 载入条款与元数据
2. `node_rule_match`    第 1 层：规则引擎快速筛查（确定性、零成本）
3. `node_llm_review`    第 2 层：LLM 语义研判（可降级）
4. `node_risk_aggregate` 合并两层结果 + 计算综合风险 + 映射结论
5. `node_persist_risk`  风险项落库 + 状态推进（completed / blocked）
6. END

为什么是纯顺序图:
    SPEC 4.3.1 给的四层是严格的流水线关系（筛查 → 研判 → 证据链 → 综合），
    层与层之间没有条件分支。唯一的"分支"语义是 LLM 失败降级，
    但按 SPEC 8.5 它**不改变流程走向**（任务仍要跑完，只是结果里少了 LLM 部分），
    所以用状态位 `llm_failed` 表达，而不是用条件边。

异常流向:
    本图**不会因为 LLM 故障抛异常**（node_llm_review 内部转成状态位）。
    真正的异常只来自：task_id 非法、数据库故障。
    这两类由调用方捕获后按 SPEC 11.1 处理。
'''

import logging

from langgraph.constants import END
from langgraph.graph import StateGraph

from processor.review_processor.base import setup_logging
from processor.review_processor.nodes.node_llm_review import NodeLlmReview
from processor.review_processor.nodes.node_load_clause import NodeLoadClause
from processor.review_processor.nodes.node_persist_risk import NodePersistRisk
from processor.review_processor.nodes.node_risk_aggregate import NodeRiskAggregate
from processor.review_processor.nodes.node_rule_match import NodeRuleMatch
from processor.review_processor.state import ReviewGraphState, get_default_state


class ContractReviewWorkflow:
    """审查子图工作流（② review_processor）"""

    def __init__(self, provider=None):
        """
        初始化工作流对象
        :param provider: LLM Provider 覆盖项。默认 None 表示按 .env 配置决定（真实/mock）；
                        传入可注入替身，用于自测与其他需要确定性结果的场景。
        """
        self._provider = provider
        self._compiled_graph = None

    @property
    def graph(self):
        """
        延迟加载（懒加载）：只在第一次使用工作流对象的时候编译图
        :return: 编译后的图
        """
        if self._compiled_graph is None:
            self._compiled_graph = self.build_graph()
        return self._compiled_graph

    def build_graph(self):
        """
        构建审查工作流图
        :return: 编译后的图
        """
        # 1.初始化图（工作流）
        graph = StateGraph(ReviewGraphState)

        # 2.在图中注册节点
        graph.add_node('node_load_clause', NodeLoadClause())  # 载入条款与元数据
        graph.add_node('node_rule_match', NodeRuleMatch())  # 第 1 层：规则引擎
        graph.add_node('node_llm_review', NodeLlmReview(provider=self._provider))  # 第 2 层：LLM 语义研判
        graph.add_node('node_risk_aggregate', NodeRiskAggregate())  # 第 3~4 层：证据链 + 综合风险
        graph.add_node('node_persist_risk', NodePersistRisk())  # 落库与状态推进

        # 3.设置入口节点
        graph.set_entry_point('node_load_clause')

        # 4.注册顺序边
        graph.add_edge('node_load_clause', 'node_rule_match')
        graph.add_edge('node_rule_match', 'node_llm_review')
        graph.add_edge('node_llm_review', 'node_risk_aggregate')
        graph.add_edge('node_risk_aggregate', 'node_persist_risk')
        graph.add_edge('node_persist_risk', END)

        # 5.编译工作流
        return graph.compile()

    def run(self, state: ReviewGraphState, stream: bool = False):
        """
        执行审查工作流
        :param state: 初始状态（只需 task_id）
        :param stream: True 时以流式返回每个节点的执行结果
        :return: 最终状态，或流式迭代器
        """
        if stream:
            return self.graph.stream(state)
        return self.graph.invoke(state)


if __name__ == '__main__':
    # 激活全局日志
    setup_logging()

    import json
    import tempfile
    from pathlib import Path

    import docx
    from sqlalchemy import select

    from config_bag.llm_config import get_llm_config
    from db.models import Clause, ContractMeta, ContractTask, RiskItem, TaskEventLog
    from db.seed import seed_all
    from db.session import get_engine, init_db, init_engine, reset_engine, session_scope
    from processor.parse_processor.main_graph import ContractParseWorkflow
    from processor.parse_processor.state import get_default_state as get_parse_default_state
    from processor.review_processor.nodes.node_rule_match import reset_rule_cache

    passed = 0

    class _StubProvider:
        """
        自测用的 LLM 替身。

        **为什么必须注入替身**：
        本项目的 .env 可能已经配置了真实 LLM（SiliconFlow / DeepSeek 等）。
        若自测直接走真实链路，会有三个问题：
          1. 每次跑自测都要烧额度、依赖网络，CI 里不可用；
          2. 模型输出不确定，断言会随机失败；
          3. 真实模型确实可能对"合规合同"给出合理的新发现，
             导致"合规=低风险"这类断言失效。
        所以图级自测统一注入本替身，把确定性交还给测试。
        真实 LLM 链路由 service 层与人工验收来验证。
        """

        is_mock = True
        name = 'stub'

        def __init__(self):
            self.call_count = 0

        def chat(self, messages, temperature=None, json_mode=False) -> str:
            self.call_count += 1
            # 固定返回空结果，让综合风险完全由规则层决定
            return '{"items": []}'

    stub_provider = _StubProvider()

    # ---------- 准备：临时库 + 规则种子 ----------
    work_dir = Path(tempfile.mkdtemp(prefix='review_graph_test_'))
    reset_engine()
    init_engine(f'sqlite:///{work_dir.as_posix()}/review_graph.db')
    init_db()
    seed_all()
    reset_rule_cache()

    def _make_risk_docx(name: str) -> Path:
        """
        造一份「有风险」的采购合同（对应 PRD 11.1 的场景）：
        知识产权归乙方 + 货到即付全款无验收 + 违约责任不对等
        """
        file_path = work_dir / 'incoming' / name
        file_path.parent.mkdir(parents=True, exist_ok=True)
        document = docx.Document()
        for text in [
            '软件采购合同',
            '合同编号：CG-2026-0918-001',
            '甲方：某某科技有限公司',
            '乙方：某某软件股份有限公司\n统一社会信用代码：91310000MA1FL1234X',
            '第一条 合同金额：人民币 1,280,000.00 元（大写：壹佰贰拾捌万元整）。',
            '第二条 付款方式：货到即付全款，未设验收条件。',
            '第三条 知识产权：本合同软件的知识产权归乙方所有。',
            '第四条 违约责任：甲方承担全部损失，赔偿责任不设赔偿上限。',
            '第五条 履行期限：自2026年1月1日起至2026年12月31日止。',
        ]:
            document.add_paragraph(text)
        document.save(str(file_path))
        return file_path

    def _make_clean_docx(name: str) -> Path:
        """造一份「基本合规」的合同：付款有验收前置、保密写明期限、管辖在甲方所在地"""
        file_path = work_dir / 'incoming' / name
        file_path.parent.mkdir(parents=True, exist_ok=True)
        document = docx.Document()
        for text in [
            '技术服务合同',
            '合同编号：JS-2026-0001',
            '甲方：某某科技有限公司',
            '乙方：某某技术服务有限公司\n统一社会信用代码：91310000MA1FL5678Y',
            '第一条 合同金额：人民币 500,000.00 元。',
            '第二条 付款方式：甲方应在验收合格后 15 个工作日内支付合同款项。',
            '第三条 保密义务：双方的保密义务自本合同生效之日起至终止后 5 年内持续有效。',
            '第四条 争议解决：协商不成的，提交甲方所在地人民法院管辖。',
        ]:
            document.add_paragraph(text)
        document.save(str(file_path))
        return file_path

    def _create_task(name: str, biz_type: str) -> int:
        """建一条任务，返回主键"""
        with session_scope() as session:
            session.add(ContractTask(contract_name=name, biz_type=biz_type, source_type='upload'))
        with session_scope() as session:
            return session.scalars(
                select(ContractTask).where(ContractTask.contract_name == name)
            ).first().id

    def _run_full_chain(task_id: int, file_path: Path, title: str) -> list:
        """跑完整链路：① 解析子图 → ② 审查子图，返回审查侧的节点执行顺序"""
        parse_state = get_parse_default_state()
        parse_state.update({
            'task_id': str(task_id),
            'source_type': 'upload',
            'import_file_path': str(file_path),
            'file_dir': str(work_dir / f'task-{task_id}'),
            'file_title': title,
        })
        ContractParseWorkflow().run(parse_state)

        review_state = get_default_state()
        review_state['task_id'] = str(task_id)

        executed = []
        for event in ContractReviewWorkflow(provider=stub_provider).run(review_state, stream=True):
            for node_name in event.keys():
                executed.append(node_name)
        return executed

    # ---------- 一、高风险合同：端到端 ①→② ----------
    logging.getLogger().info('=' * 16 + ' 一、高风险合同（①→②全链路）' + '=' * 16)

    risk_task_id = _create_task('软件采购合同-高风险', '采购')
    executed = _run_full_chain(risk_task_id, _make_risk_docx('软件采购合同.docx'), '软件采购合同')

    expected_order = [
        'node_load_clause', 'node_rule_match', 'node_llm_review',
        'node_risk_aggregate', 'node_persist_risk',
    ]
    assert executed == expected_order, \
        f'审查节点顺序不符\n  期望: {expected_order}\n  实际: {executed}'

    with session_scope() as session:
        task = session.get(ContractTask, risk_task_id)
        assert task.review_status == 'completed', f'应 completed，实际 {task.review_status}'
        assert task.risk_level == 'high', f'综合风险应为 high，实际 {task.risk_level}'

        # 元数据应被解析出来（R-01 才不会误报）
        meta = session.scalars(
            select(ContractMeta).where(ContractMeta.task_id == risk_task_id)
        ).first()
        assert meta and meta.party_b, '解析阶段应抽出相对方名称'
        assert meta.credit_code, '解析阶段应抽出统一社会信用代码'

        rows = session.scalars(select(RiskItem).where(RiskItem.task_id == risk_task_id)).all()
        risk_by_dimension = {}
        for row in rows:
            risk_by_dimension.setdefault(row.dimension, []).append(row)

        # R-04：付款无验收前置
        assert 'payment' in risk_by_dimension, \
            f'应检出付款维度风险，实际维度 {list(risk_by_dimension)}'
        assert any(r.risk_level == 'high' for r in risk_by_dimension['payment']), \
            '付款无验收前置应为高风险'
        # R-08：IP 归属不利
        assert 'ip' in risk_by_dimension, '应检出知识产权风险'
        assert any(r.risk_level == 'high' for r in risk_by_dimension['ip']), \
            'IP 归属不利应为高风险'
        # R-02：违约责任不对等
        assert 'breach' in risk_by_dimension, '应检出违约责任风险'

        # R-01 不应误报（主体信息齐全）
        assert 'subject' not in risk_by_dimension or \
            all(r.source == 'llm' for r in risk_by_dimension.get('subject', [])), \
            '主体信息齐全时 R-01 不应以规则形式命中'

        # 风险项必须能指回条款（证据链）
        linked = [r for r in rows if r.clause_id is not None]
        assert linked, '风险项应挂到具体条款上'
        linked_clause = session.get(Clause, linked[0].clause_id)
        assert linked_clause is not None and linked_clause.page_no >= 0, '条款定位信息应可用'

        # 事件链：PARSE_OK → REVIEW_OK
        events = [e.event for e in session.scalars(
            select(TaskEventLog).where(TaskEventLog.task_id == risk_task_id)
        ).all()]
        assert 'PARSE_OK' in events, '缺少 PARSE_OK 事件'
        assert 'REVIEW_OK' in events, '缺少 REVIEW_OK 事件'

    passed += 1
    logging.getLogger().info(
        f'  ✅ 全链路走通，综合风险 high，检出维度：{list(risk_by_dimension)}'
    )

    # ---------- 二、低风险合同：不应误报 ------
    logging.getLogger().info('=' * 16 + ' 二、合规合同（正向用例）' + '=' * 16)

    clean_task_id = _create_task('技术服务合同-合规', '服务')
    _run_full_chain(clean_task_id, _make_clean_docx('技术服务合同.docx'), '技术服务合同')

    with session_scope() as session:
        task = session.get(ContractTask, clean_task_id)
        assert task.review_status == 'completed', f'应 completed，实际 {task.review_status}'

        rows = session.scalars(select(RiskItem).where(RiskItem.task_id == clean_task_id)).all()
        rule_rows = [r for r in rows if r.source == 'rule']
        assert rule_rows == [], \
            f'合规合同不应产出规则类风险，实际命中 {[(r.dimension, r.risk_level) for r in rule_rows]}'

        # 注入的替身固定返回空结果 → 综合风险完全由规则层决定，应为 low
        assert task.risk_level == 'low', \
            f'规则层无命中时综合风险应为 low，实际 {task.risk_level}'
    passed += 1
    logging.getLogger().info('  ✅ 合规合同 → 规则层零误报，综合风险 low')

    # ---------- 三、降级提示语（注入 mock 替身验证）----------
    logging.getLogger().info('=' * 16 + ' 三、LLM 降级提示 ' + '=' * 16)

    with session_scope() as session:
        event = session.scalars(
            select(TaskEventLog).where(
                TaskEventLog.task_id == clean_task_id,
                TaskEventLog.event == 'REVIEW_OK',
            )
        ).first()
        detail = json.loads(event.detail_json)
        assert detail['llm_is_mock'] is True, '注入的是 mock 替身，应标记 llm_is_mock'
        assert '不可用于生产' in detail['review_note'], \
            f"降级提示应写入事件快照，实际 {detail['review_note']!r}"
        assert detail['llm_failed'] is False, 'mock 是主动降级，不应标记失败'
    passed += 1
    logging.getLogger().info('  ✅ 降级提示写入事件快照（llm_is_mock=True、含"不可用于生产"）')

    # 真实 LLM 是否已配置，仅作提示，不影响上面用例的确定性
    if get_llm_config().is_configured:
        logging.getLogger().info(
            f'  ℹ️ 检测到已配置真实 LLM（{get_llm_config().base_url} / '
            f'{get_llm_config().model}）；本自测已注入替身，未消耗额度'
        )

    # ---------- 四、invoke 模式 ----------
    logging.getLogger().info('=' * 16 + ' 四、invoke 模式 ' + '=' * 16)

    invoke_task_id = _create_task('invoke 模式', '采购')
    parse_state = get_parse_default_state()
    parse_state.update({
        'task_id': str(invoke_task_id),
        'source_type': 'upload',
        'import_file_path': str(_make_risk_docx('invoke合同.docx')),
        'file_dir': str(work_dir / f'task-{invoke_task_id}'),
        'file_title': 'invoke合同',
    })
    ContractParseWorkflow().run(parse_state)

    review_state = get_default_state()
    review_state['task_id'] = str(invoke_task_id)
    final_state = ContractReviewWorkflow(provider=stub_provider).run(review_state)

    assert final_state['risk_items'], 'invoke 模式下应有风险项'
    assert final_state['overall_risk_level'] == 'high'
    assert final_state['review_conclusion'], '应产出审查结论'
    passed += 1
    logging.getLogger().info(
        f"  ✅ invoke 模式正常：{len(final_state['risk_items'])} 条风险 / "
        f"{final_state['overall_risk_level']} / {final_state['review_conclusion']}"
    )

    reset_engine()
    logging.getLogger().info(f'review_processor/main_graph 自测全部通过 ✅ 共 {passed} 个用例')
